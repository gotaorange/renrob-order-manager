"""Synthetic XLS-only checks. Never reads private customer templates or data."""
import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest

import xlwt
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location("pi_converter", Path(__file__).resolve().parents[1] / "tools/convert_pi_template.py")
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)


class PiConversionTest(unittest.TestCase):
    def fixture(self, path, unknown_formula=False):
        book = xlwt.Workbook()
        sheet = book.add_sheet("PI")
        style = xlwt.easyxf('font: name Arial, bold on, height 240; borders: bottom thin; alignment: wrap on;', num_format_str='#,##0.00')
        sheet.write_merge(2, 2, 0, 3, 'Synthetic old buyer', style)
        sheet.write(11, 5, xlwt.Formula('D12*E12'), style)
        sheet.write(25, 1, 'Synthetic seller')
        sheet.write(27, 0, 'Synthetic terms')
        sheet.write(29, 0, 'Synthetic remaining terms')
        sheet.col(1).width = 27 * 256
        sheet.row(2).height = 44 * 20
        sheet.portrait = True
        sheet.paper_size_code = 9
        sheet.left_margin = 0.4
        sheet.right_margin = 0.4
        sheet.print_scaling = 90
        sheet.print_centered_horz = True
        # Stray formatting must not create hundreds of thousands of XLSX cells.
        sheet.write(1500, 200, '', style)
        if unknown_formula:
            sheet.write(0, 6, xlwt.Formula('1+1'))
        book.save(str(path))

    def test_preserves_format_and_print_settings_clears_order_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'source.xls', Path(directory) / 'template.xlsx'
            self.fixture(source)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            converter.convert(source, output)
            self.assertEqual(before, hashlib.sha256(source.read_bytes()).hexdigest())
            sheet = load_workbook(output).worksheets[0]
            self.assertEqual(sheet.title, 'PI')
            self.assertIn('A3:D3', [str(item) for item in sheet.merged_cells.ranges])
            self.assertEqual(sheet['A3'].value, None)
            self.assertEqual(sheet['B26'].value, None)
            self.assertEqual(sheet['A28'].value, None)
            self.assertEqual(sheet.column_dimensions['B'].width, 27)
            self.assertEqual(sheet.row_dimensions[3].height, 44)
            self.assertEqual(sheet['F12'].font.name, 'Arial')
            self.assertTrue(sheet['F12'].font.bold)
            self.assertEqual(sheet['F12'].border.bottom.style, 'thin')
            self.assertEqual(sheet['F12'].number_format, '#,##0.00')
            self.assertEqual(sheet['F12'].value, '=D12*E12')
            self.assertEqual(sheet['F23'].value, '=SUM(F20:F22)')
            self.assertEqual(sheet.page_setup.orientation, 'portrait')
            self.assertEqual(sheet.page_setup.scale, 90)
            self.assertEqual(sheet.page_margins.left, 0.4)
            self.assertLessEqual(sheet.max_row, 37)
            self.assertLessEqual(sheet.max_column, 6)
            self.assertLess(output.stat().st_size, 100000)
            self.assertTrue(sheet.print_options.horizontalCentered)
            self.assertEqual(output.stat().st_mode & 0o777, 0o600)
            with self.assertRaisesRegex(ValueError, 'already exists'):
                converter.convert(source, output)

    def test_does_not_silently_flatten_unknown_formulas(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'source.xls', Path(directory) / 'template.xlsx'
            self.fixture(source, unknown_formula=True)
            with self.assertRaisesRegex(ValueError, 'outside the supported PI layout'):
                converter.convert(source, output)
            self.assertFalse(output.exists())


if __name__ == '__main__':
    unittest.main(verbosity=2)
