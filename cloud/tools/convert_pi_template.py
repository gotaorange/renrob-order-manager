#!/usr/bin/env python3
"""One-time private PI template conversion. Requires xlrd and openpyxl.

Usage: python tools/convert_pi_template.py source.xls private/pi-template.xlsx
The source is never modified. Customer-specific input cells are cleared; the
cloud generator fills them from the selected order and private PI settings.
This intentionally accepts the existing one-sheet, six-colour PI layout only.
Unknown formula locations or unsupported embedded drawing objects fail closed.
"""
from __future__ import annotations

import argparse
import io
from pathlib import Path
import re
import struct

import xlrd
from xlrd.compdoc import CompDoc
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.workbook.properties import CalcProperties


FORMULAS = {**{f"F{row}": f"D{row}*E{row}" for row in range(12, 18)},
            "D20": "SUM(D12:D17)", "F20": "SUM(F12:F17)+F19",
            "F21": "ROUND((F20+F22)*0.05,0)", "F23": "SUM(F20:F22)"}
INPUTS = ("A3", "F5", "F6", "A11", "B12", "B18", "B19", "F19", "F22", "B26", "A28", "A30")
BORDER_STYLES = [None, "thin", "medium", "dashed", "dotted", "thick", "double", "hair",
                 "mediumDashed", "dashDot", "mediumDashDot", "dashDotDot", "mediumDashDotDot", "slantDashDot"]
FILL_STYLES = [None, "solid", "mediumGray", "darkGray", "lightGray", "darkHorizontal", "darkVertical",
               "darkDown", "darkUp", "darkGrid", "darkTrellis", "lightHorizontal", "lightVertical",
               "lightDown", "lightUp", "lightGrid", "lightTrellis", "gray125", "gray0625"]


def records(stream: bytes, start=0):
    offset = start
    while offset + 4 <= len(stream):
        record, length = struct.unpack_from("<HH", stream, offset)
        offset += 4
        if offset + length > len(stream):
            raise ValueError("Truncated XLS record")
        yield record, stream[offset:offset + length]
        offset += length
        if record == 0x000A:
            break


def sheet_records(raw: bytes):
    if raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        container = CompDoc(raw)
        stream = container.get_named_stream("Workbook") or container.get_named_stream("Book")
    else:
        stream = raw
    if not stream:
        raise ValueError("Missing XLS workbook stream")
    offsets = [struct.unpack_from("<I", data)[0] for record, data in records(stream) if record == 0x0085]
    if len(offsets) != 1:
        raise ValueError("PI conversion requires exactly one worksheet")
    return list(records(stream, offsets[0]))


def color(book, index):
    rgb = book.colour_map.get(index)
    return Color(rgb="FF" + "".join(f"{part:02X}" for part in rgb)) if rgb else Color(auto=True)


def style(book, xf):
    font = book.font_list[xf.font_index]
    border = xf.border
    underline = {1: "single", 2: "double", 33: "singleAccounting", 34: "doubleAccounting"}.get(font.underline_type)
    sides = {name: Side(style=BORDER_STYLES[getattr(border, f"{name}_line_style")],
                        color=color(book, getattr(border, f"{name}_colour_index")))
             for name in ("left", "right", "top", "bottom", "diag")}
    rotation = xf.alignment.rotation
    return {
        "font": Font(name=font.name, size=font.height / 20, bold=bool(font.bold), italic=bool(font.italic),
                     strike=bool(font.struck_out), underline=underline, color=color(book, font.colour_index),
                     vertAlign={1: "superscript", 2: "subscript"}.get(font.escapement),
                     family=font.family, charset=font.character_set),
        "fill": PatternFill(patternType=FILL_STYLES[xf.background.fill_pattern],
                            fgColor=color(book, xf.background.pattern_colour_index),
                            bgColor=color(book, xf.background.background_colour_index)),
        "border": Border(left=sides["left"], right=sides["right"], top=sides["top"], bottom=sides["bottom"],
                         diagonal=sides["diag"], diagonalUp=bool(border.diag_up), diagonalDown=bool(border.diag_down)),
        "alignment": Alignment(horizontal={0: "general", 1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify", 6: "centerContinuous", 7: "distributed"}.get(xf.alignment.hor_align, "general"),
                               vertical={0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}.get(xf.alignment.vert_align, "bottom"),
                               textRotation=rotation, wrapText=bool(xf.alignment.text_wrapped),
                               shrinkToFit=bool(xf.alignment.shrink_to_fit), indent=xf.alignment.indent_level),
        "number_format": book.format_map[xf.format_key].format_str,
        "protection": Protection(locked=bool(xf.protection.cell_locked), hidden=bool(xf.protection.formula_hidden)),
    }


def header_footer(target, data):
    if len(data) < 3:
        return
    count = struct.unpack_from("<H", data)[0]
    wide = data[2] & 1
    text = data[3:3 + count * (2 if wide else 1)].decode("utf-16le" if wide else "latin1")
    sections = re.split(r"&([LCR])", text)
    if sections[0]:
        target.center.text = sections[0]
    for index in range(1, len(sections), 2):
        getattr(target, {"L": "left", "C": "center", "R": "right"}[sections[index]]).text = sections[index + 1]


def preserve_print_settings(target, source_records):
    margins = {0x0026: "left", 0x0027: "right", 0x0028: "top", 0x0029: "bottom"}
    for record, data in source_records:
        if record in margins and len(data) >= 8:
            setattr(target.page_margins, margins[record], struct.unpack_from("<d", data)[0])
        elif record == 0x00A1 and len(data) >= 34:
            paper, scale, start, width, height, flags, hres, vres, header, footer, copies = struct.unpack_from("<8H2dH", data)
            target.page_setup.paperSize = str(paper)
            target.page_setup.scale = scale
            target.page_setup.firstPageNumber = start
            target.page_setup.fitToWidth = width
            target.page_setup.fitToHeight = height
            target.page_setup.orientation = "portrait" if flags & 2 else "landscape"
            target.page_setup.pageOrder = "overThenDown" if flags & 1 else "downThenOver"
            target.page_setup.blackAndWhite = bool(flags & 8)
            target.page_setup.draft = bool(flags & 16)
            target.page_setup.useFirstPageNumber = bool(flags & 128)
            target.page_setup.horizontalDpi = hres
            target.page_setup.verticalDpi = vres
            target.page_setup.copies = copies
            target.page_margins.header = header
            target.page_margins.footer = footer
        elif record in (0x0083, 0x0084, 0x002A, 0x002B) and len(data) >= 2:
            attribute = {0x0083: "horizontalCentered", 0x0084: "verticalCentered", 0x002A: "headings", 0x002B: "gridLines"}[record]
            setattr(target.print_options, attribute, bool(struct.unpack_from("<H", data)[0]))
        elif record == 0x0081 and len(data) >= 2:
            target.sheet_properties.pageSetUpPr.fitToPage = bool(struct.unpack_from("<H", data)[0] & 0x0100)
        elif record in (0x0014, 0x0015):
            header_footer(target.oddHeader if record == 0x0014 else target.oddFooter, data)
        elif record in (0x001B, 0x001A) and len(data) >= 2:
            count = struct.unpack_from("<H", data)[0]
            if len(data) < 2 + count * 6:
                raise ValueError("Unsupported legacy print-break record")
            breaks = target.col_breaks if record == 0x001B else target.row_breaks
            for index in range(count):
                point, first, last = struct.unpack_from("<HHH", data, 2 + index * 6)
                breaks.append(Break(id=point, min=first, max=last))


def convert(source: Path, output: Path):
    if source.suffix.lower() != ".xls" or output.suffix.lower() != ".xlsx":
        raise ValueError("Expected source.xls and output.xlsx")
    if output.exists() or output.is_symlink():
        raise ValueError("Output already exists; choose a new private filename")
    raw = source.read_bytes()
    source_records = sheet_records(raw)
    for record, data in source_records:
        if record == 0x0006:
            row, column = struct.unpack_from("<HH", data)
            address = f"{get_column_letter(column + 1)}{row + 1}"
            if address not in FORMULAS:
                raise ValueError(f"Formula at {address} is outside the supported PI layout; convert with Excel to preserve it")
        if record in (0x005D, 0x00EC, 0x00EB, 0x0207, 0x01B8, 0x01B0, 0x01BE):
            raise ValueError("Template contains drawings, hyperlinks, or conditional formatting; convert with Excel to preserve these features")
    book = xlrd.open_workbook(file_contents=raw, formatting_info=True)
    if book.nsheets != 1:
        raise ValueError("PI conversion requires exactly one worksheet")
    original = book.sheet_by_index(0)
    # BIFF used ranges often include years of stray row/column formatting. The
    # PI's visible extent is its content plus merged areas, not every formatted
    # blank cell. Materializing the full used rectangle can exceed Worker RAM.
    populated = [(row, col) for row in range(original.nrows) for col in range(original.ncols)
                 if original.cell_type(row, col) not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)]
    last_row = max([37, *[row + 1 for row, _ in populated], *[area[1] for area in original.merged_cells]])
    last_col = max([6, *[col + 1 for _, col in populated], *[area[3] for area in original.merged_cells]])
    if last_row > 200 or last_col > 30:
        raise ValueError("PI content exceeds the supported compact layout; inspect the source before converting")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = original.name
    styles = [style(book, xf) for xf in book.xf_list]
    sheet.sheet_view.showGridLines = bool(original.show_grid_lines)
    sheet.sheet_format.defaultRowHeight = original.default_row_height / 20
    if original.standardwidth:
        sheet.sheet_format.defaultColWidth = original.standardwidth / 256
    for row in range(last_row):
        for column in range(last_col):
            if row >= original.nrows or column >= original.ncols:
                continue
            cell = original.cell(row, column)
            target = sheet.cell(row + 1, column + 1)
            value = cell.value
            if cell.ctype == xlrd.XL_CELL_DATE:
                value = xlrd.xldate_as_datetime(value, book.datemode)
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                value = bool(value)
            elif cell.ctype == xlrd.XL_CELL_ERROR:
                value = xlrd.error_text_from_code[value]
            target.value = value if cell.ctype not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) else None
            for key, val in styles[cell.xf_index].items():
                setattr(target, key, val)
    for index, info in original.rowinfo_map.items():
        if index >= last_row:
            continue
        dimension = sheet.row_dimensions[index + 1]
        dimension.height = info.height / 20
        dimension.hidden = bool(info.hidden)
        dimension.outlineLevel = info.outline_level
        dimension.collapsed = bool(info.outline_group_starts_ends)
    for index, info in original.colinfo_map.items():
        if index >= last_col:
            continue
        dimension = sheet.column_dimensions[get_column_letter(index + 1)]
        dimension.width = info.width / 256
        dimension.hidden = bool(info.hidden)
        dimension.outlineLevel = info.outline_level
        dimension.collapsed = bool(info.collapsed)
    for first_row, end_row, first_col, end_col in original.merged_cells:
        sheet.merge_cells(start_row=first_row + 1, end_row=end_row, start_column=first_col + 1, end_column=end_col)
    preserve_print_settings(sheet, source_records)
    for name in book.name_obj_list:
        if name.name in ("Print_Area", "Print_Titles"):
            if not name.result or not isinstance(name.result.value, list):
                raise ValueError("Unsupported dynamic print range; convert with Excel to preserve it")
            areas = []
            for area in name.result.value:
                if area.shtxlo != 0 or area.shtxhi != 1:
                    raise ValueError("Unsupported print range outside the PI worksheet")
                rlo, rhi, clo, chi = area.rowxlo, area.rowxhi, area.colxlo, area.colxhi
                if name.name == "Print_Area":
                    areas.append(f"{get_column_letter(clo + 1)}{rlo + 1}:{get_column_letter(chi)}{rhi}")
                elif clo == 0 and chi == 256:
                    sheet.print_title_rows = f"{rlo + 1}:{rhi}"
                elif rlo == 0 and rhi == 65536:
                    sheet.print_title_cols = f"{get_column_letter(clo + 1)}:{get_column_letter(chi)}"
                else:
                    raise ValueError("Unsupported partial print titles; convert with Excel to preserve them")
            if areas:
                sheet.print_area = areas
    if not sheet.print_area:
        sheet.print_area = f"A1:{get_column_letter(last_col)}{last_row}"
    # Leave the template empty of the known order/settings fields. All generated
    # PI values and cached formulas come from the selected order on the server.
    for address in INPUTS:
        sheet[address] = None
    for row in range(12, 18):
        for column in ("C", "D", "E", "F"):
            sheet[f"{column}{row}"] = None
    for address, formula in FORMULAS.items():
        sheet[address] = "=" + formula
    workbook.calculation = CalcProperties(fullCalcOnLoad=True)
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = buffer.getvalue()
    reopened = load_workbook(io.BytesIO(result))
    checked = reopened.worksheets[0]
    assert set(str(item) for item in checked.merged_cells.ranges) == set(str(item) for item in sheet.merged_cells.ranges)
    assert all(checked[address].value == "=" + formula for address, formula in FORMULAS.items())
    assert all(checked[address].value is None for address in INPUTS)
    output.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive create avoids replacing an existing template accidentally.
    with output.open("xb") as handle:
        handle.write(result)
    output.chmod(0o600)
    print(f"Converted one PI worksheet; {len(original.merged_cells)} merged ranges preserved. Source unchanged.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    convert(args.source, args.output)
